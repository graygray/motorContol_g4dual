#!/usr/bin/env python3
# Copyright 2026 Gray Lin
# SPDX-License-Identifier: MIT
"""Generate matching editable PPTX and vector PDF. Run with python-pptx/reportlab installed."""
from pathlib import Path
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor

OUT = Path(__file__).resolve().parent
W, H = 13.333333, 7.5
BG, PANEL, WHITE, MUTED = '101D2D', '192C40', 'F4F7FA', 'B4C4D5'
TEAL, ORANGE, BLUE = '51D5BF', 'FFBE75', '82B5FF'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Motor Test Patterns | Usage & Meaning'
prs.core_properties.subject = 'motor_control_g4dual v3 operator guide'
prs.core_properties.author = 'Motor Control Project'
pdf = canvas.Canvas(str(OUT / 'test_patterns.pdf'), pagesize=(W*72,H*72))
pdf.setTitle(prs.core_properties.title)
slide = None
page = 0

def color(c): return RGBColor.from_string(c)
def box(x,y,w,h,fill=PANEL,oval=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.OVAL if oval else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill); sh.line.fill.background()
    pdf.setFillColor(HexColor('#'+fill))
    if oval: pdf.ellipse(x*72,(H-y-h)*72,(x+w)*72,(H-y)*72,fill=1,stroke=0)
    else: pdf.rect(x*72,(H-y-h)*72,w*72,h*72,fill=1,stroke=0)

def text(x,y,w,lines,size=20,ink=WHITE,bold=False,mono=False):
    if isinstance(lines,str): lines=lines.split('\n')
    leading=size*1.3
    height=(len(lines)*leading+4)/72
    assert x>=0 and y>=0 and x+w <= W+0.01 and y+height <= H, (page,lines)
    tf=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(height)).text_frame
    tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=False
    font='Courier New' if mono else 'Arial'
    pdf_font='Courier' if mono else ('Helvetica-Bold' if bold else 'Helvetica')
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=color(ink)
        p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(leading)
        pdf.setFont(pdf_font,size); pdf.setFillColor(HexColor('#'+ink))
        assert pdf.stringWidth(line,pdf_font,size) <= w*72+2, (page,line,'text too wide')
        pdf.drawString(x*72,(H-y)*72-size-i*leading,line)

def line(x1,y1,x2,y2,ink=TEAL,width=3):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=color(ink); sh.line.width=Pt(width)
    pdf.setStrokeColor(HexColor('#'+ink)); pdf.setLineWidth(width)
    pdf.line(x1*72,(H-y1)*72,x2*72,(H-y2)*72)

def arrow(x1,y1,x2,y2,ink=TEAL,width=3):
    line(x1,y1,x2,y2,ink,width)
    a=math.atan2(y2-y1,x2-x1)
    for d in [-0.48,0.48]:
        line(x2,y2,x2-0.18*math.cos(a+d),y2-0.18*math.sin(a+d),ink,width)

def curve(points,ink=TEAL,width=3):
    for a,b in zip(points,points[1:]): line(*a,*b,ink,width)

def begin(title,section='OPERATOR GUIDE',subtitle=None):
    global slide,page
    if page: pdf.showPage()
    page+=1; slide=prs.slides.add_slide(prs.slide_layouts[6])
    box(0,0,W,H,BG); box(.45,.42,.12,.24,TEAL)
    text(.72,.39,11,section.upper(),11,TEAL,True)
    text(.65,.93,12,title,32,WHITE,True)
    if subtitle: text(.68,1.64,12,subtitle,16,MUTED)
    text(.65,7.08,10,'motor_control_g4dual  /  TEST PATTERNS v3',10,MUTED)
    text(12.05,7.05,.65,f'{page:02d}',13,TEAL,True)
    # Keynote rejects the generated notes-slide parts. Keep source references
    # in docs/README.md; omit notes to retain editable slides and diagrams.

def command(cmd,y=2.13):
    box(.65,y,12,.66)
    text(.9,y+.16,11.5,cmd,20,TEAL,mono=True)

def points(items,x=7.05,y=3.05,w=5.6,size=19,gap=.89):
    for i,(head,body) in enumerate(items):
        text(x,y+i*gap,w,head,size,WHITE,True)
        text(x,y+i*gap+.36,w,body,size-3,MUTED)

def tag(x,y,label,ink=TEAL): text(x,y,5.8,label,14,ink,True)
def stop(x,y): box(x-.06,y-.06,.12,.12,ORANGE,True)

begin('Motor test patterns','FIELD GUIDE / VERSION 3','Usage, motion meaning and result interpretation')
text(.7,2.55,7,['One node. Five patterns.', 'A repeatable way to inspect', 'motor and platform behavior.'],29,WHITE,True)
text(.72,4.47,7,['Straight  /  Rotate  /  Circle  /  Square  /  S',
    'Commands: aic mc tp <pattern>', 'Editable guide + printable reference'],18,MUTED)
# Curvature-based S illustration: matches the implemented reference profile.
sx,sy=0.,0.; p=[(sx,sy)]
for i in range(300):
    d=(i+.5)*4/300; heading=4/(2*math.pi)*(1-math.cos(2*math.pi*d/4))
    sx+=4/300*math.cos(heading); sy+=4/300*math.sin(heading); p.append((sx,sy))
sp=[(8.2+x*1.28,5.9-y*1.28) for x,y in p]
curve(sp,TEAL,5); arrow(*sp[-8],*sp[-1],TEAL,5); stop(*sp[0])
text(8.2,6.35,4.3,'Ideal S reference - not test data',12,MUTED)

begin('Choose the question before the pattern','01 / PATTERN MAP','All commands execute inside the existing motor_control_node.')
rows=[('straight','Forward + reverse','Wheel balance; reverse response'),
      ('rotate','Left + right in place','Turn symmetry; angle repeatability'),
      ('circle','Constant-radius lap','Wheel-speed ratio; radial drift'),
      ('square','4 lines + 4 stopped turns','Distance/angle error accumulation'),
      ('s','Smooth left/right bend','Transition response; path correction')]
for i,(name,motion,meaning) in enumerate(rows):
    yy=2.35+i*.72; box(.65,yy,12,.61)
    text(.9,yy+.15,2,name,19,TEAL,True,True)
    text(3.0,yy+.15,3.5,motion,17)
    text(6.75,yy+.15,5.5,meaning,17,MUTED)
text(.72,6.28,12,'Default repetitions: 3 pairs for straight/rotate; 3 complete shapes for circle/square/S.',16,ORANGE)

begin('Start a test from standstill','02 / OPERATING SEQUENCE','Build/install the updated node and deploy the updated AICamera helper on the robot.')
box(.65,2.22,7.2,2.08)
text(.9,2.43,6.7,['ros2 run motor_control_g4dual motor_control_node '+chr(92),
    '  --ros-args --params-file /path/to/motor_control.yaml '+chr(92),
    '  -p enable_can:=true'],14,TEAL,mono=True)
text(.9,3.63,6.7,'One shell command; replace the parameter-file path.',14,MUTED)
text(.72,4.6,7.1,['1  Stop periodic cmd_vel publishers.',
    '2  aic mc en   ->   aic mc e diag',
    '3  Wait for fresh feedback + standstill.',
    '4  aic mc tp straight'],19,WHITE)
points([('Monitor in another terminal','aic mc tp status'),
        ('Required before starting','CAN on; motors enabled; no fault latch.'),
        ('Previous command must expire','Default cmd_vel timeout: 500 ms.')],8.2,2.3,4.4,18,.92)
text(8.2,5.2,4.4,['RPM resolution: 1.0 RPM/unit.',
    'Confirmed for this project.', 'YAML, launch and CLI agree.',
    'Other projects are unchanged.'],15,ORANGE)

begin('Straight: compare forward and reverse','03 / STRAIGHT','One repetition = forward 1 m, stop, reverse 1 m, stop.')
command('aic mc tp straight')
arrow(1.35,3.8,5.8,3.8); arrow(5.8,4.65,1.35,4.65,BLUE)
stop(1.35,3.8); stop(5.8,3.8); stop(1.35,4.65)
tag(2.55,3.28,'FORWARD  1 m'); tag(2.55,5.05,'REVERSE  1 m',BLUE)
text(1.3,5.85,5.3,['Same physical line; paths offset here for clarity.',
    'Amber dots indicate a standstill check.'],12,MUTED)
points([('What it tells you','Left/right wheel balance and reverse response.'),
        ('Default motion','0.1 m/s; 3 forward/reverse pairs.'),
        ('How it finishes','Distance projected onto each segment heading.'),
        ('Interpretation','Return error is measured, not corrected.')],y=3.15,gap=.81)

begin('Rotate: isolate turning behavior','04 / ROTATE','One repetition = left 90 degrees, stop, right 90 degrees, stop.')
command('aic mc tp rotate')
center=(3.5,4.65)
curve([(center[0]+1.25*math.cos(t),center[1]-1.25*math.sin(t)) for t in [i*math.pi/2/70 for i in range(71)]])
arrow(3.69,3.41,3.5,3.4)
arrow(3.5,4.65,4.68,4.65,BLUE); arrow(3.5,4.65,3.5,3.48)
box(3.29,4.44,.42,.42,WHITE,True)
tag(1.3,5.78,'LEFT 90 deg  /  RIGHT 90 deg')
points([('What it tells you','Turning symmetry and repeatable angles.'),
        ('Default motion','0.2 rad/s; 3 left/right pairs.'),
        ('Configurable angle','angle_deg supports 90, 180 or 360 degrees.'),
        ('How it finishes','Continuous yaw handles the +/-180 boundary.')],y=3.15,gap=.81)

begin('Circle: compare inner and outer wheels','05 / CIRCLE','One repetition = one full lap, then standstill. Default: counterclockwise.')
command('aic mc tp circle')
cx,cy,rr=3.6,4.55,1.35
circle=[(cx+rr*math.cos(i*2*math.pi/120),cy-rr*math.sin(i*2*math.pi/120)) for i in range(121)]
curve(circle); arrow(*circle[26],*circle[30]); stop(cx+rr,cy)
line(cx,cy,cx+rr,cy,MUTED,1); tag(2.98,4.76,'R = 1 m')
text(1.4,6.1,5.3,'Robot-centre radius; wheel paths differ.',14,MUTED)
points([('What it tells you','Wheel-speed ratio and radius consistency.'),
        ('Default motion','Radius 1 m; 3 laps; about 64 s per lap.'),
        ('Motion relationship','v = radius x angular speed; limits apply.'),
        ('How it finishes','Accumulated yaw reaches 360 deg, then stop.')],y=3.15,gap=.81)

begin('Square: expose accumulated error','06 / SQUARE','One repetition = four forward sides and four stationary 90-degree turns.')
command('aic mc tp square')
coords=[(2.0,5.65),(5.1,5.65),(5.1,3.35),(2.,3.35),(2.,5.65)]
for a,b in zip(coords,coords[1:]): arrow(*a,*b)
for pt in coords[:4]: stop(*pt)
tag(2.6,4.15,'1 m / side'); text(2.65,4.67,2.7,'8 segments',17,MUTED)
text(1.22,6.1,5.65,'Stop after every line and turn; schematic only.',14,MUTED)
points([('What it tells you','Distance, turn and stop errors accumulate.'),
        ('Default motion','1 m sides; counterclockwise; 3 squares.'),
        ('Fixed corners','Always 90 deg; angle_deg only affects rotate.'),
        ('How it finishes','Includes the last turn toward the start heading.')],y=3.15,gap=.81)

begin('Smooth S: test transition and correction','07 / S CURVE','One repetition = a continuous left/right bend; stop only at the end.')
command('aic mc tp s')
sp=[(1.45+x*1.5,6.0-y*1.1) for x,y in p]
curve(sp); arrow(*sp[-8],*sp[-1]); stop(*sp[0]); stop(*sp[-1])
text(1.38,3.12,5.5,'NO STOP AT THE MIDDLE',15,TEAL,True)
text(1.3,6.25,5.6,'Ideal reference; end is displaced from start.',14,MUTED)
points([('What it tells you','Smooth turning response and tracking error.'),
        ('Default geometry','4 m travel length; 1 m minimum nominal radius.'),
        ('Path correction','Heading + lateral gains default to 1.0.'),
        ('For motor-only response','Set both S gains to 0, then restart the node.')],y=3.15,gap=.81)

begin('Change geometry, limits and direction','08 / PARAMETERS','Edit motion_test.* in motor_control.yaml, then restart the node.')
rows=[('Parameter suffix','Default','Meaning'),
('distance_m / angle_deg','1.0 / 90.0','Straight segment / rotate angle'),
('circle_radius_m / square_side_m','1.0 / 1.0','Circle radius / square side'),
('s_length_m / s_radius_m','4.0 / 1.0','S travel length / minimum nominal radius'),
('clockwise','false','Circle/square clockwise; S right-first if true'),
('linear_speed_mps / angular_speed_radps','0.1 / 0.2','Body velocity limits'),
('linear_accel_mps2 / angular_accel_radps2','0.1 / 0.2','Body acceleration limits'),
('s_heading_gain / s_lateral_gain','1.0 / 1.0','S correction: 1/s and 1/m^2'),
('repetitions','3','Pairs or complete shapes, depending on pattern')]
for i,(a,b,c) in enumerate(rows):
    yy=2.25+i*.43
    if i%2==0: box(.65,yy,12,.42)
    text(.82,yy+.085,5.05,a,12.8,TEAL if i==0 else WHITE, i==0, i!=0)
    text(6.02,yy+.085,1.6,b,13,MUTED)
    text(7.62,yy+.085,4.9,c,13,MUTED)
text(.75,6.4,11.9,'The feature defaults to enabled. It does not start tests or enable motors automatically.',16,ORANGE)

begin('Normal commands can take over','09 / COMMAND OWNERSHIP','Tests and normal driving share the same node and motor enable state.')
for x,label,col in [(1.,'TEST RUNNING',TEAL),(5.05,'cmd_vel',BLUE),(9.1,'NORMAL DRIVE',WHITE)]:
    box(x,2.65,3.1,.85); text(x+.2,2.9,2.75,label,18,col,True)
arrow(4.17,3.08,4.88,3.08); arrow(8.22,3.08,8.95,3.08)
points([('Valid cmd_vel, including zero','Ends the test and takes control with that command.'),
        ('Normal completion','Zero speed; motors stay enabled for the next command.'),
        ('Cancel, fault or emergency stop','Motion is gated; explicitly re-enable after resolving it.')],.95,4.0,11.5,20,.83)
text(.95,6.66,11.5,'Pause periodic cmd_vel publishers while testing, or their next message will interrupt the test.',14,ORANGE)

begin('Read the result, not only the status','10 / CSV & MEANING','aic mc tp status shows state, segment, outcome reason and the CSV path.')
box(.7,2.33,5.4,3.93)
text(.96,2.58,4.95,'ILLUSTRATIVE STOPPED RESULT',13,TEAL,True)
text(.96,3.16,4.95,['Target       1.000 m', 'Measured     1.025 m', 'Difference  +0.025 m'],20,WHITE,mono=True)
text(.96,4.58,4.95,'+25 mm overshoot',26,ORANGE,True)
text(.96,5.25,4.95,['last_stopped_segment_progress', '- last_stopped_segment_target'],13,MUTED,mono=True)
points([('Check units and segment type','line / s_curve: m     turn / arc: rad'),
        ('Compare targets and feedback','Normalized wheel RPM; account for feedback delay.'),
        ('Inspect path_error_m','Lateral / radial / turn translation / S cross-track.'),
        ('CSV directory on the robot','/tmp/motor_control_tests (configurable)')],6.55,2.6,6.05,18,.87)
text(.76,6.61,12,'completed = sequence finished. Wheel odometry is not independent proof of ground-track accuracy.',15,ORANGE)

begin('If a test will not start or stops early','11 / TROUBLESHOOTING','A published command is a request. Check status for acceptance or the reason for interruption.')
rows=[('Rejected before start','Check enabled motors, CAN, fresh position + speed feedback.'),
('Recent cmd_vel / moving wheels','Stop command publishers; wait for timeout and standstill.'),
('Aborted: external cmd_vel','Normal command took control; the test will not resume.'),
('No progress / stale feedback','Inspect obstruction, wheel direction, encoders and CAN.'),
('Segment timeout','60 s per movement/wait; circle motion has a 120 s default.'),
('CSV cannot be written','Use a writable local log directory; check disk space.')]
for i,(a,b) in enumerate(rows):
    yy=2.27+i*.62; box(.65,yy,12,.53)
    text(.86,yy+.13,4.0,a,16,TEAL,True)
    text(5.05,yy+.13,7.3,b,15,MUTED)
text(.8,6.3,12,['Feedback timeout: 500 ms. Stall timeout: 5 s. Standstill: both wheels <= 0.5 RPM for 0.5 s.',
    'Cancel requests the existing ramp stop; an aborted state does not confirm physical standstill.'],14,ORANGE)

begin('Commands to keep beside the robot','12 / QUICK REFERENCE','Use one pattern at a time. Begin with short, low-speed runs and inspect the record.')
cmds=[('aic mc tp straight','Forward / reverse'),('aic mc tp rotate','Left / right rotation'),
      ('aic mc tp circle','Full circle'),('aic mc tp square','Stopped-corner square'),
      ('aic mc tp s','Smooth S'),('aic mc tp status','Monitor outcome + CSV'),
      ('aic mc tp cancel','Cancel and gate motion'),('aic mc tp help','Show command help')]
for i,(a,b) in enumerate(cmds):
    col=i//4; row=i%4; x=.75+col*6.35; y=2.3+row*.86
    text(x,y,5.9,a,21,TEAL,mono=True); text(x,y+.4,5.9,b,16,MUTED)
box(.65,6.08,12,.62)
text(.9,6.24,11.5,'Observe -> compare -> tune -> repeat. Validate physical accuracy using external measurements.',16,WHITE)

prs.save(str(OUT / 'test_patterns.pptx'))
pdf.save()
print(f'Created {page} slides: test_patterns.pptx and test_patterns.pdf')
